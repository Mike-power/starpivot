# -*- coding: utf-8 -*-
"""从 论文初稿-全文v1.0.md 生成《航天器工程》双盲投稿稿 docx。
三线表 / GB/T 7714 参考文献 / 图随文走 / 去内部注记。"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = Path("/Users/maoli/Documents/个人/能力提升/能力提升与寻找机会/星枢-星地协同智能体/docs")
SRC = BASE / "论文初稿-全文v1.0.md"
OUT = BASE / "投稿附件-航天器工程" / "投稿稿_基于输出置信度的星地动态分层推理方法_双盲v1.0.docx"

FIGS = {
    1: ("图1_星地协同分层推理系统架构_2026.10.2.png", "图 1  系统总体架构与置信度路由决策流"),
    2: ("图2_过境窗口时间线示意_2026.10.2.png", "图 2  过境窗口时间线示意（b-EDF 排队与应急抢占可行性检查）"),
    3: ("fig3_param_heatmap.png", "图 3  链路受限扫描双热力图（100 任务 × 5 种子）"),
    4: ("fig4_theta_curve.png", "图 4  能力阈值敏感性（天花板线性律与协同增益）"),
    5: ("fig5_tau_pareto.png", "图 5  τ-成功率-误收率多目标帕累托平面（k=10 / k=5）"),
    6: ("fig6_energy_bars.png", "图 6  能耗效率对比（单成功能耗五方案）"),
    7: ("fig7_constellation_scaling.png", "图 7  星座规模 scaling（成功率与带宽双轴）"),
    8: ("L7_准入策略双轴占优_2026.10.2.png", None),  # 题注已在正文
}
TAB_TITLES = {
    1: "算法在线开销逐项核算（n 为队列长度，N 为任务总量）",
    2: "三方案基线对比（TLE 场景，无排队语义）",
    3: "主实验：真实过境下的排队感知调度（六种策略）",
    4: "参数扫描节选：过境周期 × 窗口时长",
    5: "能力阈值 θ 扫描（天花板线性律）",
    6: "应急重规划矩阵节选",
    7: "置信度路由消融节选（k × 校准 × τ）",
    8: "能耗效率：五方案单成功能耗对比",
    9: "多星组网 scaling（相位偏移合成星座）",
    10: "工作流补偿 × 组网联合矩阵节选",
    11: "推理服务耗时三口径对比",
    12: "三杠杆叠加节选（N × 口径）",
    13: "准入策略对照（L7）",
}

REFS = [
 "SHI Y, ZHU J, JIANG C, 等. Satellite edge artificial intelligence with large models: architectures and technologies[J]. Science China Information Sciences, 2025.",
 "SANG H, ZHANG L, CHEN T, 等. Onboard deployment of remote sensing foundation models: a comprehensive review of architecture, optimization, and hardware[J]. Remote Sensing, 2026, 18(2): 298.",
 "LIN J, TANG J, TANG H, 等. AWQ: activation-aware weight quantization for LLM compression and acceleration[C]//Proceedings of MLSys. 2023.",
 "LEVIATHAN Y, KALMAN M, MATIAS Y. Fast inference from transformers via speculative decoding[C]//Proceedings of ICML. 2023.",
 "ZHANG S, WU W, LI L, 等. Communication-efficient collaborative LLM inference over LEO satellite networks[EB/OL]. 2026.",
 "LI Z, YANG J, ZHANG Y, 等. Enabling near-realtime remote sensing via satellite-ground collaboration of large vision-language models[EB/OL]. 2026.",
 "FAN H, LONG J, LIU L, 等. Dynamic digital twin and online scheduling for contact window resources in satellite network[J]. IEEE Transactions on Industrial Informatics, 2022.",
 "TANG Q, XIE R, FANG Z, 等. Joint service deployment and task scheduling for satellite edge computing: a two-timescale hierarchical approach[EB/OL]. 2024.",
 "LUO H, LIU Y, ZHANG R, 等. Toward edge general intelligence with multiple-large language model (Multi-LLM): architecture, trust, and orchestration[J]. IEEE Transactions, 2025.",
 "BEHERA A P, CHAMPATI J P, MORABITO R, 等. Towards efficient multi-LLM inference: characterization and analysis of LLM routing and hierarchical techniques[EB/OL]. 2025.",
 "XU D, LUO Y, PEI S, 等. Space-native hardware: a vision for a new architecture for on-satellite LLM inference[EB/OL]. 2025.",
 "NEVO M, FATKIEV O, ZAIDMAN G, 等. Payload operation using on-board vision language model[EB/OL]. 2025.",
 "WANG F, YAO H, HE W, 等. Time-sensitive scheduling mechanism based on end-to-end collaborative latency tolerance for low-earth-orbit satellite networks[EB/OL]. 2023.",
 "LIU Y, JIANG L, QI Q, 等. Online computation offloading for collaborative space/aerial-aided edge computing toward 6G system[EB/OL]. 2023.",
 "KOLAWOLE S, DENNIS D, TALWALKAR A, 等. Agreement-based cascading for efficient inference[EB/OL]. 2024.",
 "ZHANG T, MEHRADFAR A, DIMITRIADIS D, 等. Leveraging uncertainty estimation for efficient LLM routing[EB/OL]. 2025.",
 "MOSLEM Y, KELLEHER J D. Dynamic model routing and cascading for efficient LLM inference: a survey[EB/OL]. 2026.",
 "GAO X, WANG Y, LIU B, 等. Agentic satellite-augmented low-altitude economy and terrestrial networks: a survey on generative approaches[J]. IEEE Communications Surveys & Tutorials, 2026.",
 "JAVAID S, KHALIL R A, SAEED N, 等. Leveraging large language models for integrated satellite-aerial-terrestrial networks: recent advances and future directions[EB/OL]. 2024.",
 "国星宇航 Qwen3 大模型完成太空在轨部署[N]. 科技日报, 2026-01-01.",
 "CHEN Z, 等. 星地大小模型协同推理方法研究[Z]. 上海市自然科学基金项目(25ZR1401016), 2026.",
 "商业航天智能化落地的大语言模型轻量化路径研究[J]. 中国航天, 2026(8).",
 "CHEN J, 等. Design of remote sensing satellite task scheduling algorithm with large language model[J]. Engineering, 2025.",
 "CHEN J, 等. Improving scheduling model for remote sensing satellites with large language model[J]. IEEE Transactions on Geoscience and Remote Sensing, 2026.",
 "CHEN J, 等. Validating LLM-designed scheduling for remote sensing satellites[C]//Proceedings of IEEE Congress on Evolutionary Computation. 2025.",
]

# ---------- 字体与样式工具 ----------

def set_run(run, size=10.5, bold=False, east="宋体", ascii_f="Times New Roman", color=None):
    run.font.name = ascii_f
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.append(rf)
    rf.set(qn('w:eastAsia'), east)
    if color:
        run.font.color.rgb = color

def para_fmt(p, indent_chars=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.3,
             space_before=0, space_after=0):
    pf = p.paragraph_format
    pf.alignment = align
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = spacing
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if indent_chars:
        ind = p._element.get_or_add_pPr().find(qn('w:ind'))
        if ind is None:
            ind = OxmlElement('w:ind'); p._element.get_or_add_pPr().append(ind)
        ind.set(qn('w:firstLineChars'), str(indent_chars * 100))
        ind.set(qn('w:firstLine'), str(int(indent_chars * 210)))

def add_rich(p, text, size=10.5, bold=False, east="宋体"):
    """**bold** 解析；`code` 去反引号。"""
    for i, seg in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not seg:
            continue
        seg = seg.replace('`', '')
        set_run(p.add_run(seg), size=size, bold=bold or (i % 2 == 1), east=east)

def add_caption(doc, text, kind="表"):
    p = doc.add_paragraph()
    para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.15,
             space_before=6, space_after=3)
    add_rich(p, text, size=10.5, bold=False, east="黑体")
    for r in p.runs:
        r.font.bold = True
    return p

def three_line(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for tag, sz in (("top", 12), ("bottom", 12)):
        el = OxmlElement(f'w:{tag}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
        el.set(qn('w:color'), '000000'); borders.append(el)
    for tag in ("left", "right", "insideH", "insideV"):
        el = OxmlElement(f'w:{tag}')
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); borders.append(el)
    tblPr.append(borders)

def header_bottom_border(row):
    for cell in row.cells:
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        bot = OxmlElement('w:bottom')
        bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '6')
        bot.set(qn('w:color'), '000000')
        b.append(bot); tcPr.append(b)

def add_table(doc, rows):
    n_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    three_line(table)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.1)
            val = val.replace("**", "").replace("`", "")
            set_run(p.add_run(val), size=9, bold=(ri == 0), east="宋体")
        if ri == 0:
            header_bottom_border(table.rows[0])
    return table

def add_figure(doc, n):
    fname, cap = FIGS[n]
    path = BASE / "figures" / fname
    p = doc.add_paragraph()
    para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0,
             space_before=6, space_after=3)
    p.add_run().add_picture(str(path), width=Cm(14))
    if cap:
        add_caption(doc, cap, kind="图")

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    sizes = {1: 14, 2: 12}
    para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.LEFT,
             spacing=1.3, space_before=10 if level == 1 else 6, space_after=6)
    set_run(p.add_run(text), size=sizes[level], bold=True, east="黑体")
    return p

def add_footer_pagenum(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = "1"
    r.append(t); fld.append(r)
    p._element.append(fld)

# ---------- 解析全文 ----------

def clean(s):
    s = re.sub(r'（文件：docs/figures/[^）]*）', '', s)
    s = re.sub(r'（全文初稿[^）]*）', '', s)
    s = s.strip()
    # 内部版本/里程碑标记清理
    s = re.sub(r'（v[\d.]+(/v[\d.]+)?，[^）]*）', '', s)   # （v0.8/v0.8c，2026.10.2 回填）
    s = re.sub(r'（v[\d.]+系列实验）', '', s)
    s = re.sub(r'（v[\d.]+(/v[\d.]+)?）', '', s)
    s = re.sub(r'（L\d，[^）]*）', '', s)                   # （L7，2026.10.2 回填）
    s = re.sub(r'（L\d）', '', s)
    s = re.sub(r'（\d{4}\.\d{1,2}\.\d{1,2} 成稿）', '', s)
    s = s.replace('（待回填）', '')
    s = re.sub(r'，锚定 2026-10-02', '', s)
    s = re.sub(r'（摘要 v[\d.]+[^）]*）', '', s)
    # 表号全局重排：3.5 复杂度表成为表 1，实验章表 1–12 顺延为 2–13
    s = re.sub(r'表 (\d+)', lambda m: f"表 {int(m.group(1)) + 1}", s)
    return s

def main():
    raw = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    add_footer_pagenum(doc)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    i = 0
    state = {"title_done": False, "fig_done": set(), "tab_n": 0,
             "in_refs": False, "skip_until_h2": False}
    first_h1 = True

    while i < len(raw):
        ln = raw[i]
        s = ln.strip()

        # 参考文献区接管
        if state["in_refs"]:
            i += 1
            continue
        if s.startswith("## 参考文献"):
            add_heading(doc, "参考文献", 1)
            for n, ref in enumerate(REFS, 1):
                p = doc.add_paragraph()
                para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, spacing=1.15)
                pf = p.paragraph_format
                pf.left_indent = Pt(21); pf.first_line_indent = Pt(-21)
                set_run(p.add_run(f"[{n}] {ref}"), size=9)
            state["in_refs"] = True
            i += 1
            continue

        # 跳内部待办节
        if s.startswith("## 待办"):
            state["skip_until_h2"] = True; i += 1; continue
        if state["skip_until_h2"]:
            if s.startswith("## ") and not s.startswith("###"):
                state["skip_until_h2"] = False
            else:
                i += 1; continue

        if not s or s == "---":
            i += 1; continue
        if s.startswith(">"):
            i += 1; continue

        # 标题（H1）
        if s.startswith("# ") and not state["title_done"]:
            title = clean(s[2:])
            p = doc.add_paragraph()
            para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.3,
                     space_before=6, space_after=6)
            set_run(p.add_run(title), size=16, bold=True, east="黑体")
            state["title_done"] = True
            i += 1; continue

        # 章节标题
        if s.startswith("### "):
            add_heading(doc, clean(s[4:]), 2); i += 1; continue
        if s.startswith("## "):
            if state.get("skip48"):
                state["skip48"] = False
            h = clean(s[3:])
            add_heading(doc, h, 1)
            first_h1 = False
            if h.startswith("4.8"):
                p = doc.add_paragraph()
                para_fmt(p)
                add_rich(p, "星上压缩流水线（任务域蒸馏 → AWQ 4bit 量化 → Jetson Orin 部署）的实测结果，"
                            "包括压缩比、单 token 延迟、显存占用与相对地面大模型的端到端加速比，"
                            "将在终稿补充；届时 4.1 节概率契约的 σ(·) 参数与 4.11 节 TTFT/TPOT 延迟常数"
                            "将同步替换为实测值。")
                state["skip48"] = True
            i += 1; continue

        if state.get("skip48"):
            if s.startswith("## "):
                state["skip48"] = False
            else:
                i += 1; continue

        # 表格
        if s.startswith("|"):
            rows = []
            while i < len(raw) and raw[i].strip().startswith("|"):
                cells = [c.strip() for c in raw[i].strip().strip("|").split("|")]
                if not re.match(r'^[\s:\-]+$', "".join(cells)):
                    rows.append(cells)
                i += 1
            state["tab_n"] += 1
            n = state["tab_n"]
            add_caption(doc, f"表 {n}  {TAB_TITLES.get(n, '')}")
            add_table(doc, rows)
            continue

        # 图表题注行（**图 8** ...）
        m = re.match(r'^\*\*图 (\d)\*\*\s*(.*)', s)
        if m:
            n = int(m.group(1))
            if n not in state["fig_done"]:
                add_figure(doc, n)
                state["fig_done"].add(n)
            add_caption(doc, clean(s))
            i += 1; continue
        m = re.match(r'^\*\*表 (\d+)\*\*\s*(.*)', s)
        if m:
            add_caption(doc, clean(s))
            i += 1; continue

        # 英文题名 / Abstract / Keywords
        m = re.match(r'^\*\*Title\*\*[:：]?\s*(.*)', s)
        if m:
            p = doc.add_paragraph()
            para_fmt(p, indent_chars=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.2,
                     space_after=6)
            set_run(p.add_run(clean(m.group(1))), size=12, bold=True)
            i += 1; continue
        m = re.match(r'^\*\*Abstract\*\*[:：]?\s*(.*)', s)
        if m:
            p = doc.add_paragraph()
            para_fmt(p, indent_chars=0)
            set_run(p.add_run("Abstract: "), size=10.5, bold=True)
            add_rich(p, clean(m.group(1)))
            i += 1; continue
        m = re.match(r'^\*\*Keywords\*\*[:：]?\s*(.*)', s)
        if m:
            p = doc.add_paragraph()
            para_fmt(p, indent_chars=0)
            set_run(p.add_run("Keywords: "), size=10.5, bold=True)
            add_rich(p, clean(m.group(1)))
            i += 1; continue
        m = re.match(r'^\*\*关键词\*\*[:：]?\s*(.*)', s)
        if m:
            p = doc.add_paragraph()
            para_fmt(p, indent_chars=0)
            set_run(p.add_run("关键词: "), size=10.5, bold=True, east="黑体")
            add_rich(p, clean(m.group(1)))
            i += 1; continue

        # 普通段落
        p = doc.add_paragraph()
        para_fmt(p)
        add_rich(p, clean(s))
        # 图随文走：段中首次引用 图 N 且尚未插图 → 段后插图（含题注）
        for n in range(1, 9):
            if n in state["fig_done"]:
                continue
            if re.search(rf'图 {n}(?!\d)', clean(s)):
                add_figure(doc, n)
                state["fig_done"].add(n)
        i += 1

    missing = [n for n in range(1, 9) if n not in state["fig_done"]]
    print("未插图:", missing if missing else "无")
    print("表格数:", state["tab_n"])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("saved:", OUT)

if __name__ == "__main__":
    main()
